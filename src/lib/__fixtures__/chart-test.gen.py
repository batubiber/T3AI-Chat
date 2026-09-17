#!/usr/bin/env python3
"""chart-test.pptx üreteci — grafik parse testlerinin fixture'ı.

Elle yazılmış XML'in kaçıracağı namespace/yapı detaylarını yakalayabilmek için
fixture GERÇEK Office XML'i olmalı. Bu script onu yeniden üretilebilir kılar
(air-gapped ortamda kimse "bu binary nereden geldi" diye sormasın).

Kullanım:
    pip install --target ./pylibs python-pptx
    PYTHONPATH=./pylibs python3 chart-test.gen.py

Kapsadığı durumlar:
  slayt 1 — sütun grafik, 2 seri, başlıklı + konuşmacı notu  (c:cat / c:val)
  slayt 2 — pasta grafik, tek seri, BAŞLIKSIZ                (c:cat / c:val)
  slayt 3 — dağılım grafiği                                   (c:xVal / c:yVal)
"""
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE

prs = Presentation()
blank = prs.slide_layouts[6]

# Slayt 1: sütun grafik, çok seri, başlıklı, konuşmacı notlu
s1 = prs.slides.add_slide(blank)
cd = CategoryChartData()
cd.categories = ["Q1", "Q2", "Q3"]
cd.add_series("2024 Satış", (120.0, 95.5, 210.0))
cd.add_series("2025 Satış", (140.0, 130.0, 260.0))
chart = s1.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), cd
).chart
chart.has_title = True
chart.chart_title.text_frame.text = "Çeyreklik Satışlar"
s1.notes_slide.notes_text_frame.text = "Bu grafik yıllık büyümeyi gösteriyor."

# Slayt 2: pasta grafik, tek seri, başlıksız
s2 = prs.slides.add_slide(blank)
cd2 = CategoryChartData()
cd2.categories = ["Ankara", "İstanbul", "İzmir"]
cd2.add_series("Pay", (45.0, 35.0, 20.0))
s2.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(1), Inches(1), Inches(6), Inches(4), cd2)

# Slayt 3: dağılım — c:cat/c:val YOK, c:xVal/c:yVal var (ayrı kod dalı)
s3 = prs.slides.add_slide(blank)
xy = XyChartData()
ser = xy.add_series("Ölçüm")
ser.add_data_point(1.0, 2.5)
ser.add_data_point(2.0, 4.1)
s3.shapes.add_chart(
    XL_CHART_TYPE.XY_SCATTER, Inches(1), Inches(1), Inches(6), Inches(4), xy
)

prs.save("chart-test.pptx")
print("chart-test.pptx yazıldı")
